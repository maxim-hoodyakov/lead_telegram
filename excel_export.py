"""
Работа с Excel-файлом, в котором хранятся лиды и переписка.

Структура листа:
- Столбец A: ID лида.
- Столбец B: информация о клиенте (имя, username, телефон).
- Столбец C: краткая сводка о лиде от ИИ.
- Столбец D: название группы, где найден лид.
- Столбец E и далее: сообщения в диалоге (чередуются "я" / "лид").
"""

from __future__ import annotations

import os
from typing import Dict, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

import config

# Кэш соответствия ID лида и номера строки в Excel
_lead_row_cache: Dict[int, int] = {}


def clear_all_data() -> None:
    """
    Полностью очистить данные из Excel: оставить только заголовки (строка 1).
    Очищает кэш _lead_row_cache.
    """
    global _lead_row_cache
    _lead_row_cache.clear()
    if not os.path.exists(config.EXCEL_PATH):
        return
    wb = load_workbook(config.EXCEL_PATH)
    if config.EXCEL_SHEET not in wb.sheetnames:
        wb.close()
        return
    ws = wb[config.EXCEL_SHEET]
    while ws.max_row > 1:
        ws.delete_rows(2)
    wb.save(config.EXCEL_PATH)


def _load_workbook() -> Workbook:
    """
    Открыть существующий Excel-файл или создать новый с заголовками.
    """
    if os.path.exists(config.EXCEL_PATH):
        wb = load_workbook(config.EXCEL_PATH)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = config.EXCEL_SHEET
        # Заголовки
        ws.append(["ID", "Клиент", "Сводка о лиде", "Группа", "Диалог"])
        wb.save(config.EXCEL_PATH)
        return wb

    # Убедимся, что нужный лист существует
    if config.EXCEL_SHEET in wb.sheetnames:
        return wb

    ws = wb.create_sheet(title=config.EXCEL_SHEET)
    ws.append(["ID", "Клиент", "Сводка о лиде", "Группа", "Диалог"])
    wb.save(config.EXCEL_PATH)
    return wb


def _get_sheet(wb: Workbook) -> Worksheet:
    return wb[config.EXCEL_SHEET]


def _build_cache(ws: Worksheet) -> None:
    """
    Построить кэш lead_id -> row_index по уже существующим строкам.
    """
    _lead_row_cache.clear()
    for row in ws.iter_rows(min_row=2, values_only=False):
        cell = row[0]  # столбец A
        lead_id = cell.value
        if isinstance(lead_id, int):
            _lead_row_cache[lead_id] = cell.row


def init_excel_file() -> None:
    """
    Убедиться, что Excel-файл и нужный лист существуют.
    """
    wb = _load_workbook()
    ws = _get_sheet(wb)
    _build_cache(ws)
    wb.save(config.EXCEL_PATH)


def ensure_lead_row(lead_id: int, user_info: str, group_title: str) -> int:
    """
    Гарантировать, что для лида есть строка в Excel.
    Если строки нет — создать новую. Вернуть номер строки.
    """
    wb = _load_workbook()
    ws = _get_sheet(wb)

    if not _lead_row_cache:
        _build_cache(ws)

    if lead_id in _lead_row_cache:
        return _lead_row_cache[lead_id]

    # Добавляем новую строку: A=id, B=клиент, C=сводка (пусто), D=группа
    row_index = ws.max_row + 1
    ws.cell(row=row_index, column=1, value=lead_id)
    ws.cell(row=row_index, column=2, value=user_info)
    ws.cell(row=row_index, column=3, value="")
    ws.cell(row=row_index, column=4, value=group_title)
    _lead_row_cache[lead_id] = row_index

    wb.save(config.EXCEL_PATH)
    return row_index


def append_message(lead_id: int, text: str, from_me: bool) -> None:
    """
    Добавить сообщение к существующей строке лида.
    Сообщения пишутся в первую свободную ячейку справа, начиная с столбца E (5).
    """
    wb = _load_workbook()
    ws = _get_sheet(wb)

    if not _lead_row_cache:
        _build_cache(ws)

    row_index = _lead_row_cache.get(lead_id)
    if row_index is None:
        row_index = ws.max_row + 1
        ws.cell(row=row_index, column=1, value=lead_id)
        ws.cell(row=row_index, column=2, value="")
        ws.cell(row=row_index, column=3, value="")
        ws.cell(row=row_index, column=4, value="")
        _lead_row_cache[lead_id] = row_index

    # Первая пустая колонка — с E (5)
    col = 5
    while ws.cell(row=row_index, column=col).value not in (None, ""):
        col += 1

    ws.cell(row=row_index, column=col, value=text)
    wb.save(config.EXCEL_PATH)


def update_lead_summary(lead_id: int, summary: str) -> None:
    """
    Обновить краткую информацию о лиде в столбце C (Сводка о лиде).
    """
    wb = _load_workbook()
    ws = _get_sheet(wb)

    if not _lead_row_cache:
        _build_cache(ws)

    row_index = _lead_row_cache.get(lead_id)
    if row_index is None:
        return
    ws.cell(row=row_index, column=3, value=summary)
    wb.save(config.EXCEL_PATH)


def get_lead_excel_info(lead_id: int) -> Optional[str]:
    """
    Прочитать из Excel короткую информацию о лиде (ID, Клиент, Сводка, Группа).
    Возвращает строку для уведомления или None, если строки нет.
    """
    if not os.path.exists(config.EXCEL_PATH):
        return None
    wb = _load_workbook()
    ws = _get_sheet(wb)
    if not _lead_row_cache:
        _build_cache(ws)
    row_index = _lead_row_cache.get(lead_id)
    if row_index is None:
        return None
    parts = []
    for col, label in enumerate(["ID", "Клиент", "Сводка о лиде", "Группа"], start=1):
        val = ws.cell(row=row_index, column=col).value
        if val is not None and str(val).strip():
            parts.append(f"{label}: {val}")
        elif label == "ID":
            parts.append(f"ID: {lead_id}")
    return "\n".join(parts) if parts else None

